"""Prepare complete Dataset-1 HVLM manufacturing rules, never runtime state."""
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
from openpyxl import load_workbook
from ..inputs import bind_validation,verify_binding,validate_keys
from ..timebase import to_ticks

ADAPTER='smt2020-hvlm-v1'
VALIDATOR='manufacturing-v1'
UNITS={'sec':Decimal(1),'s':Decimal(1),'min':Decimal(60),'hr':Decimal(3600),'day':Decimal(86400)}
HEADERS={
 'Toolgroups':('AREA','TOOLGROUP','NUMBER OF TOOLS','CASCADINGTOOL','BACTHINGTOOL','BATCHCRITERION','BATCHING UNIT','LOADINGTIME','LT UNITS','UNLOADINGTIME','ULT UNITS','TOOLGROUPLOCATION','DISPATCHING','Ranking 1','Ranking 2','Ranking 3','TOOL WAKE UP Ranking'),
 'PM':('PM EVENT NAME','PM EVENT VALID FOR TYPE','TYPE NAME','PM TYPE','MTBeforePM','MTBPM UNITS','TTR DISTRIBUTION','MEAN','OFFSET','TTR UNITS','FIRST ONE AT DISTRIBUTION','FOA','FOA UNITS'),
 'Breakdown':('DOWN EVENT NAME','DOWN EVENT VALID FOR TYPE','TYPE NAME','DOWN TYPE','TTF DISTRIBUTION','MTTF','MTTF UNITS','TTR DISTRIBUTION','MTTR','MTTR UNITS','FIRST ONE AT DISTRIBUTION','FOA','FOA UNITS'),
 'route':('ROUTE','STEP','STEP DESCRIPTION','AREA','TOOLGROUP','PROCESSING UNIT','PROCESSINGTIME DISTRIBUTION','MEAN','OFFSET','PT UNITS','CASCADING INTERVAL','C Units','BATCH MINIMUM','BATCH MAXIMUM','SETUP','WHEN','SETUP DISTRIBUTION','SETUP TIME','OFFSET','ST UNITS','STEP FOR LTL DEDICATION','REWORK PROBABILITY in %','R UNIT','STEP FOR REWORK','PROCESSING PROBABILITY in % (Sampling)','STEP FOR CRITICAL QUEUE TIME','CQT','CQTUNITS'),
 'Lotrelease':('PRODUCT NAME','ROUTE NAME','LOT NAME/TYPE','PRIORITY','SUPERHOTLOT','WAFERS PER LOT','START DATE','RELEASE DISTRIBUTION','RELEASE INTERVAL','R UNITS','LOTS PER RELEASE','DUE DATE','Release Scenario'),
 'Setups':('SETUP GROUP NAME','CURRENT SETUP','NEW SETUP','SETUP TIME','ST UNITS','MINMAL NUMBER OF RUNS'),
 'Transport':('FROM LOCATION','TO LOCATION','TRANSPORTTIME DISTRIBUTION','MEAN','OFFSET','TT UNITS'),
}

class SourceError(ValueError):pass

def decimal(value,where,positive=False):
    try:
        if isinstance(value,bool) or value is None:raise ValueError()
        result=Decimal(str(value))
    except (InvalidOperation,ValueError):raise SourceError(f'{where}: expected finite number') from None
    if not result.is_finite() or result<0 or (positive and result==0):raise SourceError(f'{where}: invalid numeric range')
    return result

def text(value):
    return format(value,'f').rstrip('0').rstrip('.') if '.' in format(value,'f') else format(value,'f')

def seconds(value,unit,where,positive=False):
    if unit not in UNITS:raise SourceError(f'{where}: unsupported time unit {unit!r}')
    return decimal(value,where,positive)*UNITS[unit]

def integer(value,where,positive=False):
    v=decimal(value,where,positive)
    if v!=v.to_integral_value():raise SourceError(f'{where}: expected integer')
    return int(v)

def distribution(kind,mean,offset,unit,where):
    if kind not in ('constant','uniform','exponential'):raise SourceError(f'{where}: unsupported distribution {kind!r}')
    m=seconds(mean,unit,where)
    o=seconds(0 if offset is None else offset,unit,where)
    if kind!='uniform' and o:raise SourceError(f'{where}: offset requires uniform distribution')
    if kind=='uniform' and o>m:raise SourceError(f'{where}: invalid uniform range')
    if kind=='exponential' and m==0:raise SourceError(f'{where}: exponential mean must be positive')
    return {'distribution':kind,'mean_seconds':text(m),'offset_seconds':text(o)}

def flag(value,where):
    if value is None or value=='':return False
    if isinstance(value,str) and value.upper() in ('YES','NO'):return value.upper()=='YES'
    raise SourceError(f'{where}: unsupported flag {value!r}')

def probability(value,where,default=100):
    v=decimal(default if value is None else value,where)
    if v>100:raise SourceError(f'{where}: probability outside [0,100]')
    return text(v/100)

def _rows(book,title):
    if title not in book.sheetnames:raise SourceError(f'{title}: missing sheet')
    rows=list(book[title].values)
    expected=HEADERS['route' if title.startswith('Route_') else title]
    if not rows or tuple(rows[0])!=expected:raise SourceError(f'{title} row 1: unexpected header layout')
    result=[]
    for n,row in enumerate(rows[1:],2):
        if not any(v is not None for v in row):continue
        result.append((n,row))
    return result

def _unique(records,where):
    ids=[r['id'] for r in records]
    if any(not isinstance(i,str) or not i for i in ids) or len(ids)!=len(set(ids)):
        raise SourceError(f'{where}: missing/duplicate identity')

def _source_value(v):
    if isinstance(v,datetime):return v.isoformat()
    if isinstance(v,(int,float)) and not isinstance(v,bool):return str(v)
    return v

def _components(steps):
    by={s['source_step_id']:s for s in steps};parent={k:k for k in by};edges=[]
    def find(x):
        while parent[x]!=x:x=parent[x]
        return x
    for s in steps:
        target=s['dedication_target']
        if target is None:continue
        if target not in by or by[target]['tool_group_id']!=s['tool_group_id'] or by[target]['kind']!='processing':
            raise SourceError(f"{s['source']['sheet']} row {s['source']['row']}: incompatible dedication target {target}")
        a,b=find(s['source_step_id']),find(target);parent[b]=a;edges.extend([s['source_step_id'],target])
    out={}
    for k in sorted(set(edges)):out.setdefault(find(k),[]).append(k)
    return [{'id':'dedication:'+':'.join(sorted(v)), 'source_step_ids':sorted(v)} for v in out.values()]

def prepare_manufacturing(path: Path) -> dict:
    path=Path(path);book=load_workbook(path,read_only=True,data_only=True)
    try:return _prepare(book,hashlib.sha256(path.read_bytes()).hexdigest())
    finally:book.close()

def _prepare(book,source_hash):
    source_tables={}
    def rows(title):
        result=_rows(book,title)
        source_tables[title]={'headers':list(HEADERS['route' if title.startswith('Route_') else title]),
                              'rows':[{'row':n,'values':[_source_value(v) for v in r]} for n,r in result]}
        return result
    groups=[]
    for n,r in rows('Toolgroups'):
        where=f'Toolgroups row {n}';cascade=flag(r[3],where);batch=flag(r[4],where)
        if cascade and batch:raise SourceError(f'{where}: simultaneous batch/cascade unsupported')
        if batch and (r[5]!='Same Product and Same Step' or r[6]!='Wafer'):raise SourceError(f'{where}: unsupported batch rule')
        if tuple(r[13:16])!=('Highest Lotpriority','Least Setuptime','FIFO'):raise SourceError(f'{where}: unsupported ranking')
        if r[12] not in ('SuperHotLotFIRST_and_Reservation','Setupavoidancerule'):raise SourceError(f'{where}: unsupported dispatch')
        groups.append({'id':r[1],'area':r[0],'source_count':integer(r[2],where,True),'cascading':cascade,'batching':batch,
                       'loading_us':to_ticks(seconds(r[7],r[8],where)),'unloading_us':to_ticks(seconds(r[9],r[10],where)),
                       'source_dispatch':r[12],'dispatch_policy':'priority_setup_fifo_no_lookahead_v1',
                       'rankings':list(r[13:16]),'source_location':r[11]})
    _unique(groups,'Toolgroups');bygroup={g['id']:g for g in groups}
    setups=[]
    for n,r in rows('Setups'):
        where=f'Setups row {n}'
        if not r[2]:raise SourceError(f'{where}: missing new setup')
        setups.append({'id':f'setup:{n}','setup_group':r[0],'current':r[1],'new':r[2],
                       'duration_us':to_ticks(seconds(r[3],r[4],where)),
                       'minimum_runs':integer(r[5],where) if r[5] is not None else 0})
    keys=[(s['current'],s['new']) for s in setups]
    if len(keys)!=len(set(keys)):raise SourceError('Setups: ambiguous transitions')
    routes=[]
    for product in ('Product_3','Product_4'):
        title='Route_'+product;steps=[]
        source_rows=rows(title)
        expected_count=583 if product=='Product_3' else 343
        sequence=[str(integer(r[1],f'{title} row {n} STEP',True)) for n,r in source_rows]
        if sequence!=[str(i) for i in range(1,expected_count+1)]:
            raise SourceError(f'{title}: expected complete source sequence 1..{expected_count}')
        for n,r in source_rows:
            where=f'{title} row {n}';sid=str(integer(r[1],where,True));group=bygroup.get(r[4])
            if not group or r[0]!=title or r[3]!=group['area']:raise SourceError(f'{where}: missing/inconsistent route/tool group')
            if r[5] not in ('Lot','Wafer','Batch'):raise SourceError(f'{where}: unsupported processing unit')
            process=distribution(r[6],r[7],r[8],r[9],where+' processing')
            abstraction=r[4]=='Delay_32'
            interval=to_ticks(seconds(r[10],r[11],where,True)) if r[10] is not None else None
            if group['cascading']!=(interval is not None):raise SourceError(f'{where}: cascade interval/flag mismatch')
            bmin=bmax=None
            if r[5]=='Batch':
                bmin=integer(r[12],where,True);bmax=integer(r[13],where,True)
                if bmin>bmax or not group['batching']:raise SourceError(f'{where}: invalid batch range/flag')
            elif r[12] is not None or r[13] is not None or group['batching']:raise SourceError(f'{where}: unsupported batch semantics')
            setup=None
            if r[14]:
                if r[15] not in ('need','always'):raise SourceError(f'{where} WHEN: unsupported value')
                if not r[16] and any(v is not None for v in r[17:20]):
                    raise SourceError(f'{where}: partial setup fallback definition')
                fallback=distribution(r[16],r[17],r[18],r[19],where+' setup') if r[16] else None
                setup={'id':r[14],'when':r[15],'fallback':fallback}
            elif any(v is not None for v in r[15:20]):raise SourceError(f'{where}: setup parameters without setup')
            rework=None
            if r[21] is not None or r[23] is not None:
                if not isinstance(r[22],str) or r[22].lower()!='lot':raise SourceError(f'{where}: unsupported rework unit')
                rework={'probability':probability(r[21],where+' rework',0),'target_step_id':str(integer(r[23],where+' rework',True))}
            cqt=None
            if any(v is not None for v in r[25:28]):
                cqt={'target_step_id':str(integer(r[25],where+' CQT',True)),
                     'threshold_us':to_ticks(seconds(r[26],r[27],where+' CQT'))}
            steps.append({'id':title+':'+sid,'source_step_id':sid,'kind':'transport_abstraction' if abstraction else 'processing',
                          'tool_group_id':r[4],'description':r[2],'processing_unit':r[5].lower(),'processing':process,
                          'cascade_interval_us':interval,'batch_min':bmin,'batch_max':bmax,'setup':setup,
                          'sampling_probability':probability(r[24],where+' sampling'),
                          'dedication_target':str(integer(r[20],where+' dedication',True)) if r[20] is not None else None,
                          'rework':rework,'cqt':cqt,'source':{'sheet':title,'row':n}})
        _unique(steps,title)
        by={s['source_step_id']:s for s in steps}
        for i,s in enumerate(steps):
            if s['kind']=='transport_abstraction':
                s['replacement']={'policy':'explicit_amhs_v1','next_executable_step_id':next((t['source_step_id'] for t in steps[i+1:] if t['kind']=='processing'),None)}
            for name in ('rework','cqt'):
                rule=s[name]
                if rule and rule['target_step_id'] not in by:raise SourceError(f'{title} row {s["source"]["row"]}: invalid {name} target')
            if s['cqt'] and by[s['cqt']['target_step_id']]['sampling_probability']!='1':raise SourceError(f'{title}: sampled CQT target unsupported')
        routes.append({'id':title,'product_id':product,'steps':steps,'dedication_components':_components(steps)})
    required=sorted({s['tool_group_id'] for r in routes for s in r['steps'] if s['kind']=='processing'})
    groups=[g for g in groups if g['id'] in required]
    for group in groups:
        group_steps=[s for r in routes for s in r['steps'] if s['tool_group_id']==group['id'] and s['kind']=='processing']
        names=sorted({s['setup']['id'] for s in group_steps if s['setup']})
        for s in group_steps:
            setup=s['setup']
            if not setup:continue
            for current in names:
                if current==setup['id'] and setup['when']=='need':continue
                matches=[t for t in setups if t['new']==setup['id'] and t['current']==current]
                if not matches:matches=[t for t in setups if t['new']==setup['id'] and t['current'] is None]
                if not matches and setup['fallback'] is None:raise SourceError(f'{s["source"]}: unresolved setup {current}->{setup["id"]}')
        group['initial_setup']=names[0] if names else None
    maintenance=[]
    for n,r in rows('PM'):
        where=f'PM row {n}'
        if r[1]!='toolgroup' or r[2] not in required:raise SourceError(f'{where}: unsupported scope')
        if r[3] not in ('time based (calendar)','counter based'):raise SourceError(f'{where}: unsupported type')
        counter=r[3]=='counter based'
        if counter:
            if r[5]!='Wafer' or r[12]!='Wafer' or r[10]!='constant':raise SourceError(f'{where}: unsupported counter unit/distribution')
            interval=integer(r[4],where,True);first={'distribution':'constant','value':text(decimal(r[11],where))}
        else:
            interval=to_ticks(seconds(r[4],r[5],where,True));first=distribution(r[10],r[11],None,r[12],where+' first')
        duration=distribution(r[6],r[7],r[8],r[9],where+' duration')
        if not counter and duration['distribution']=='exponential':
            raise SourceError(f'{where}: calendar PM duration must be bounded')
        if not counter and to_ticks(Decimal(duration['mean_seconds'])+Decimal(duration['offset_seconds']))>interval:
            raise SourceError(f'{where}: PM duration exceeds calendar interval')
        maintenance.append({'id':r[0],'scope':'toolgroup','scope_id':r[2],'type':'counter' if counter else 'calendar',
                            'interval':interval,'first':first,'duration':duration})
    breakdowns=[];areas={g['area'] for g in groups}
    for n,r in rows('Breakdown'):
        where=f'Breakdown row {n}'
        if r[1]!='area' or r[2] not in areas:raise SourceError(f'{where}: unsupported scope')
        if r[3]!='time based (calendar)':raise SourceError(f'{where}: unsupported breakdown type')
        breakdowns.append({'id':r[0],'scope':'area','scope_id':r[2],
                           'ttf':distribution(r[4],r[5],None,r[6],where+' TTF'),
                           'duration':distribution(r[7],r[8],None,r[9],where+' duration'),
                           'first':distribution(r[10],r[11],None,r[12],where+' first')})
    releases=[];epoch=datetime(2018,1,1)
    for n,r in rows('Lotrelease'):
        where=f'Lotrelease row {n}'
        if r[0] not in ('Product_3','Product_4') or r[1]!='Route_'+r[0]:raise SourceError(f'{where}: unsupported product/route')
        if r[7]!='constant':raise SourceError(f'{where}: unsupported release distribution')
        if not isinstance(r[6],datetime) or not isinstance(r[11],datetime) or r[6].tzinfo or r[11].tzinfo:raise SourceError(f'{where}: invalid source date')
        delta=r[6]-epoch;due=r[11]-r[6]
        if delta.total_seconds()<0 or due.total_seconds()<0:raise SourceError(f'{where}: negative release/due allowance')
        if r[4] not in ('yes','no'):raise SourceError(f'{where}: invalid SUPERHOTLOT')
        wafer=integer(r[5],where,True)
        if wafer!=25:raise SourceError(f'{where}: standard cohort requires 25-wafer LOTs')
        releases.append({'id':r[2],'product_id':r[0],'route_id':r[1],'priority':integer(r[3],where),
                         'super_hot':r[4]=='yes','wafer_count':wafer,'first_at_us':delta.days*86400000000+delta.seconds*1000000+delta.microseconds,
                         'interval_seconds':text(seconds(r[8],r[9],where+' interval',True)),
                         'lots_per_release':integer(r[10],where,True),'due_allowance_us':due.days*86400000000+due.seconds*1000000+due.microseconds})
    for records,label in [(maintenance,'PM'),(breakdowns,'Breakdown'),(releases,'Lotrelease')]:_unique(records,label)
    for n,r in rows('Transport'):
        if r[0]!='Fab' or r[1]!='Fab':raise SourceError(f'Transport row {n}: unsupported abstraction')
        distribution(r[2],r[3],r[4],r[5],f'Transport row {n}')
    # This matrix is explanatory source metadata; Setups is the executable authority.
    source_tables['Setup_Matrix_Implant_Gas']={'rows':[{'row':n,'values':[_source_value(v) for v in r]} for n,r in enumerate(book['Setup_Matrix_Implant_Gas'].values,1) if any(v is not None for v in r)]}
    model={'id':'smt2020-dataset1-hvlm','products':[{'id':p,'route_id':'Route_'+p} for p in ('Product_3','Product_4')],
           'routes':routes,'tool_groups':sorted(groups,key=lambda g:g['id']),'setup_transitions':setups,
           'maintenance':sorted(maintenance,key=lambda r:r['id']),'breakdowns':sorted(breakdowns,key=lambda r:r['id']),
           'releases':sorted(releases,key=lambda r:r['id']), 'required_tool_group_ids':required,
           'policies':{'uniform':'mean_plus_minus_offset_v1','abstract_transport':'explicit_amhs_v1','baseline_maintenance':'source_causes_v1'}}
    artifact={'schema_version':1,'model':model,'provenance':{'source_sha256':source_hash,'dataset':'SMT2020 Dataset 1 HVLM',
              'adapter_version':ADAPTER,'source_tables':source_tables}}
    return bind_validation(artifact,{'tool_groups':required},VALIDATOR)

def validate_manufacturing(artifact: dict) -> dict:
    validate_keys(artifact,{'schema_version','model','provenance','validation'},'manufacturing')
    if type(artifact.get('schema_version')) is not int or artifact['schema_version']!=1:raise SourceError('unsupported manufacturing schema')
    if artifact.get('provenance',{}).get('adapter_version')!=ADAPTER:raise SourceError('unsupported manufacturing adapter version')
    requirements={'tool_groups':artifact['model']['required_tool_group_ids']}
    if not verify_binding(artifact,requirements,VALIDATOR):raise SourceError('manufacturing content/validation identity mismatch; reprepare source')
    return artifact
