"""Independent SMAT physical adapter; no Dataset-4 manufacturing inference."""
import hashlib
from pathlib import Path
from openpyxl import load_workbook
from .smt import decimal, integer, text, _source_value, validate_manufacturing
from ..geometry import spatial_um, chord_length
from ..inputs import bind_validation, canonical_digest, validate_keys, verify_binding
from ..topology import compile_traffic, index

ADAPTER = 'smat2022-chord-center-v1'
VALIDATOR = 'topology-v1'
ALIASES = {'DefMEt_FE_118': 'DefMet_FE_118'}
HEADERS = {
    'Address': ('NAME','FAB_NAME','POSITION_X','POSITION_Y','ZCU_NAME','ZCU_TYPE'),
    'ZCU': ('NAME','FAB_NAME'),
    'Rail': ('NAME','FAB_NAME','BAY_NAME','FROM_NODE','TO_NODE','MAX_SPEED','CURVE'),
    'Bay': ('NAME','FAB_NAME','TYPE','RETICLE','NEIGHBOR_BAY'),
    'Equipment': ('NAME','TOOL_GROUP','BAY','WIDTH','HEIGHT','POSITION_X','POSITION_Y'),
    'PortType': ('NAME','SIZE','INOUT_TYPE'),
    'Port': ('NAME','FAB_NAME','PORT_TYPE','EQP_NAME','RAILLINE_NAME','DISTANCE'),
    'VehicleType': ('NAME','SIZE','MAX_SPEED','ACCELERATION','DECELERATION','MIN_DISTANCE'),
    'Vehicle': ('NAME','FAB_NAME','TYPE','RAILLINE_NAME','DISTANCE'),
}


def traffic_requirements(manufacturing):
    validate_manufacturing(manufacturing)
    requirements = {}
    for group in manufacturing['model']['tool_groups']:
        if group['id'] not in manufacturing['model']['required_tool_group_ids']: continue
        steps = [s for route in manufacturing['model']['routes'] for s in route['steps']
                 if s['tool_group_id']==group['id'] and s['kind']=='processing']
        capacity = 1
        if group['batching']: capacity = max(s['batch_max']//25 for s in steps)
        elif group['cascading'] and any(s['processing_unit']=='lot' for s in steps): capacity = 2
        requirements[group['id']] = {'rack_min':capacity}
    return {'manufacturing_sha256':manufacturing['validation']['content_sha256'],
            'tool_groups':requirements, 'geometry_policy':'chord_center_full_envelope_v1'}


def prepare_topology(path: Path, manufacturing: dict, overrides: dict) -> dict:
    requirements = traffic_requirements(manufacturing)
    validate_keys(overrides, {'stocker_capacity','rack_capacities','footprints',
                             'fleet_count','vehicle_positions','holding_pitch_mm',
                             'curve_cap_mm_s','pickup_seconds','drop_seconds'}, 'physical overrides')
    if 'stocker_capacity' not in overrides:
        raise ValueError('physical profile requires explicit stocker_capacity; PortType.SIZE is not carrier capacity')
    stocker_capacity = integer(overrides['stocker_capacity'], 'stocker_capacity', True)
    path = Path(path); book = load_workbook(path, read_only=True, data_only=True)
    source = {}; tables = {}
    try:
        if set(book.sheetnames) != set(HEADERS): raise ValueError('unexpected SMAT source sheets')
        for name, expected in HEADERS.items():
            values = list(book[name].values)
            if not values or tuple(values[0]) != expected: raise ValueError(f'{name} row 1: unexpected header layout')
            rows = [(n,r) for n,r in enumerate(values[1:],2) if any(v is not None for v in r)]
            identities = [r[0] for n,r in rows]
            if any(not isinstance(i,str) or not i for i in identities) or len(identities)!=len(set(identities)):
                raise ValueError(f'{name}: missing/duplicate identity')
            tables[name] = rows
            source[name] = {'headers':list(expected), 'rows':[{'row':n,'values':[_source_value(v) for v in r]} for n,r in rows]}
    finally: book.close()
    def dim(value, where):
        decimal(value, where, True); result = spatial_um(value)
        if result<=0: raise ValueError(f'{where}: dimension quantizes to zero')
        return result
    def coord(value, where):
        try: return spatial_um(value)
        except ValueError as error: raise ValueError(f'{where}: {error}') from None
    bays = [{'id':r[0], 'type':r[2], 'source_neighbors':r[4]} for n,r in tables['Bay']]
    bay_ids = set(index(bays,'bays'))
    nodes = [{'id':r[0], 'x_um':coord(r[2],f'Address row {n} POSITION_X'),
              'y_um':coord(r[3],f'Address row {n} POSITION_Y'),
              'zcu_id':r[4] if r[4] not in (None,'',',') else None, 'zcu_type':r[5]}
             for n,r in tables['Address']]
    by_node = index(nodes,'nodes'); rails = []
    for n,r in tables['Rail']:
        where=f'Rail row {n}'
        if r[2] not in bay_ids or r[3] not in by_node or r[4] not in by_node: raise ValueError(f'{where}: missing bay/node reference')
        if type(r[6]) not in (int,float) or r[6] not in (0,1): raise ValueError(f'{where} CURVE: unsupported classification')
        multiplier=decimal(r[5],where+' MAX_SPEED',True)
        if multiplier>1: raise ValueError(f'{where} MAX_SPEED: multiplier must be <=1')
        a,b=by_node[r[3]],by_node[r[4]]
        rails.append({'id':r[0], 'bay_id':r[2], 'from_node':r[3], 'to_node':r[4],
                      'length_um':chord_length((a['x_um'],a['y_um']),(b['x_um'],b['y_um'])),
                      'curve':bool(r[6]), 'speed_multiplier':text(multiplier), 'geometry_mode':'chord'})
    vehicle_types=[]
    for n,r in tables['VehicleType']:
        where=f'VehicleType row {n}'
        decimal(r[3],where+' ACCELERATION',True); decimal(r[4],where+' DECELERATION',True)
        vehicle_types.append({'id':r[0], 'length_um':dim(r[1],where+' SIZE'),
                              'speed_um_s':text(decimal(r[2],where+' MAX_SPEED',True)*1000),
                              'gap_um':spatial_um(decimal(r[5],where+' MIN_DISTANCE')),
                              'carrier_capacity':1})
    equipment=[]; footprints=overrides.get('footprints',{}); racks=overrides.get('rack_capacities',{})
    physical_to_manufacturing={ALIASES.get(g,g):g for g in requirements['tool_groups']}
    if len(physical_to_manufacturing)!=len(requirements['tool_groups']): raise ValueError('ambiguous physical group alias')
    for mapping,name in ((footprints,'footprints'),(racks,'rack_capacities')):
        if not isinstance(mapping,dict): raise ValueError(f'{name}: expected object')
        if set(mapping)-set(requirements['tool_groups']): raise ValueError(f'{name}: unknown ordinary tool group')
    for n,r in tables['Equipment']:
        where=f'Equipment row {n}'; physical_group=ALIASES.get(r[1],r[1])
        group=physical_to_manufacturing.get(physical_group,physical_group)
        if r[2] not in bay_ids: raise ValueError(f'{where}: missing bay reference')
        if group in requirements['tool_groups']:
            kind='ordinary'; capacity=requirements['tool_groups'][group]['rack_min']
            capacity=integer(racks.get(group,capacity),where+' rack capacity',True)
            if capacity<requirements['tool_groups'][group]['rack_min']: raise ValueError(f'{where}: rack capacity below manufacturing requirement')
        elif group in ('STOCKER','COMMIT','COMPLETE'):
            kind={'STOCKER':'stocker','COMMIT':'entry','COMPLETE':'exit'}[group]; capacity=stocker_capacity
        else: raise ValueError(f'{where}: unsupported tool group {group}')
        width,height=dim(r[3],where+' WIDTH'),dim(r[4],where+' HEIGHT')
        if group in footprints:
            profile=footprints[group]; validate_keys(profile,{'width_mm','height_mm'},where+' footprint override')
            width,height=dim(profile['width_mm'],where+' width_mm'),dim(profile['height_mm'],where+' height_mm')
        equipment.append({'id':r[0], 'tool_group_id':group, 'kind':kind, 'bay_id':r[2],
                          'physical_reference_group_id':physical_group,
                          'width_um':width, 'height_um':height, 'x_um':coord(r[5],where+' POSITION_X'),
                          'y_um':coord(r[6],where+' POSITION_Y'), 'coordinate_reference':'center', 'rack_capacity':capacity})
    by_eq=index(equipment,'equipment')
    present={e['tool_group_id'] for e in equipment if e['kind']=='ordinary'}
    if present!=set(requirements['tool_groups']): raise ValueError('physical source missing required tool groups')
    port_types=[]
    for n,r in tables['PortType']:
        if r[2] not in ('IN','OUT','BOTH'): raise ValueError(f'PortType row {n}: unsupported role')
        port_types.append({'id':r[0],'size_um':dim(r[1],f'PortType row {n} SIZE'),'role':r[2]})
    by_port_type=index(port_types,'port types'); docks=[]
    for n,r in tables['Port']:
        where=f'Port row {n}'
        if r[2] not in by_port_type: raise ValueError(f'{where}: unknown port type')
        if r[3] in by_eq:
            eq_id=r[3]; kind=by_eq[eq_id]['kind']; capacity=1
        elif r[3] in (None,'',',') and r[2] in ('LSTB','RSTB','UTB'):
            eq_id=None; kind='track_buffer'; capacity=1
        else: raise ValueError(f'{where}: unsupported equipment/storage mapping')
        docks.append({'id':r[0], 'type_id':r[2], 'equipment_id':eq_id, 'kind':kind,
                      'rail_id':r[4], 'offset_um':spatial_um(decimal(r[5],where+' DISTANCE')),
                      'role':by_port_type[r[2]]['role'], 'carrier_capacity':capacity})
    roles={e['id']:set() for e in equipment}
    for d in docks:
        if d['equipment_id']: roles[d['equipment_id']].add(d['role'])
    for e in equipment:
        have=roles[e['id']]
        expected={'entry':'OUT','exit':'IN'}.get(e['kind'])
        if expected:
            if not (expected in have or 'BOTH' in have): raise ValueError(f"equipment {e['id']}: missing service role")
        elif not ('BOTH' in have or {'IN','OUT'}<=have): raise ValueError(f"equipment {e['id']}: missing IN/OUT service")
    vehicles=[{'id':r[0], 'type_id':r[2], 'rail_id':r[3], 'offset_um':spatial_um(decimal(r[4],f'Vehicle row {n} DISTANCE')),
               'carrier_id':None} for n,r in tables['Vehicle']]
    if 'fleet_count' in overrides:
        count=integer(overrides['fleet_count'],'fleet_count',True)
        if count>len(vehicles): raise ValueError('fleet_count exceeds supplied fleet; explicit new fleet requires generated topology')
        vehicles=sorted(vehicles,key=lambda v:v['id'])[:count]
    positions=overrides.get('vehicle_positions',{})
    if not isinstance(positions,dict) or set(positions)-{v['id'] for v in vehicles}: raise ValueError('unknown vehicle_positions')
    for v in vehicles:
        if v['id'] in positions:
            p=positions[v['id']]; validate_keys(p,{'rail_id','offset_mm'},'vehicle_positions')
            v.update(rail_id=p['rail_id'],offset_um=spatial_um(decimal(p['offset_mm'],'vehicle_positions offset_mm')))
    from ..timebase import to_ticks
    model={'id':'smat2022-lvhm-physical', 'nodes':nodes, 'rails':rails, 'bays':bays,
           'zcus':[{'id':r[0]} for n,r in tables['ZCU']], 'equipment':equipment,
           'port_types':port_types, 'docks':docks, 'vehicle_types':vehicle_types, 'vehicles':vehicles,
           'policy':{'geometry':'chord_integer_sqrt_ties_even_v1', 'coordinate_reference':'center',
                     'holding_pitch_um':dim(overrides.get('holding_pitch_mm',2000),'holding_pitch_mm'),
                     'curve_cap_um_s':text(decimal(overrides.get('curve_cap_mm_s',1000),'curve_cap_mm_s',True)*1000),
                     'pickup_us':to_ticks(decimal(overrides.get('pickup_seconds',10),'pickup_seconds',True)),
                     'drop_us':to_ticks(decimal(overrides.get('drop_seconds',10),'drop_seconds',True)),
                     'raw_dock_clamp':True}}
    if not model['policy']['pickup_us'] or not model['policy']['drop_us']: raise ValueError('handling must remain positive after quantization')
    compile_traffic(model)
    artifact={'schema_version':1,'model':model,
              'provenance':{'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                            'dataset':'SMAT2022 LVHM physical source', 'adapter_version':ADAPTER,
                            'aliases':ALIASES.copy(), 'units_interpretation':'mm',
                            'equipment_origin_interpretation':'center', 'overrides':overrides,
                            'source_tables':source}}
    return bind_validation(artifact,requirements,VALIDATOR)


def validate_topology(artifact: dict, manufacturing: dict) -> dict:
    validate_keys(artifact,{'schema_version','model','provenance','validation'},'topology')
    if type(artifact.get('schema_version')) is not int or artifact['schema_version']!=1:
        raise ValueError('unsupported topology schema')
    if artifact.get('provenance',{}).get('adapter_version')!=ADAPTER: raise ValueError('unsupported topology adapter')
    if not verify_binding(artifact,traffic_requirements(manufacturing),VALIDATOR):
        raise ValueError('topology content/validation identity mismatch; reprepare source')
    return artifact
