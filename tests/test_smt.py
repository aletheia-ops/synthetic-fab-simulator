from copy import deepcopy
import os
from pathlib import Path
import pytest
from openpyxl import load_workbook
from fab_sim.sources.smt import prepare_manufacturing,validate_manufacturing

@pytest.fixture(scope='module')
def source():
    p=os.environ.get('FAB_HVLM_WORKBOOK')
    if not p:pytest.skip('set FAB_HVLM_WORKBOOK for complete Dataset-1 source tests')
    return Path(p)

@pytest.fixture(scope='module')
def prepared(source):return prepare_manufacturing(source)

def route(a,p):return next(r for r in a['model']['routes'] if r['product_id']==p)

def test_complete_routes_and_delay_mapping(prepared):
    r3=route(prepared,'Product_3');r4=route(prepared,'Product_4')
    assert len(r3['steps'])==583 and len(r4['steps'])==343
    assert sum(s['kind']=='transport_abstraction' for s in r3['steps'])==24
    assert sum(s['kind']=='transport_abstraction' for s in r4['steps'])==10
    assert len(prepared['model']['tool_groups'])==105
    assert r3['steps'][0]['source_step_id']=='1' and r3['steps'][-1]['source_step_id']=='583'

def test_duplicate_offsets_and_source_process_values(prepared):
    s=route(prepared,'Product_3')['steps'][0]
    assert s['processing']=={'distribution':'uniform','mean_seconds':'30079.8','offset_seconds':'1503.99'}
    assert s['setup'] is None
    assert s['batch_min']==125 and s['batch_max']==150
    assert prepared['provenance']['source_sha256']=='11c5cc9f477eafcbb1d37abf2a6802eade5964a95fa1b1fd3351ddce841c1a18'

def test_complete_source_rules_and_references(prepared):
    r=route(prepared,'Product_3')
    assert len(r['dedication_components'])>0
    assert sum(s['rework'] is not None for s in r['steps'])==7
    assert sum(s['cqt'] is not None for s in r['steps'])==41
    setups=prepared['model']['setup_transitions']
    assert sum(s['minimum_runs']==7 for s in setups)==9
    assert len(prepared['model']['maintenance'])==292
    assert len(prepared['model']['breakdowns'])==11
    assert all(b['scope']=='area' for b in prepared['model']['breakdowns'])

def test_source_release_values_and_due_allowances(prepared):
    rel=prepared['model']['releases']
    assert len(rel)==5
    regular=next(r for r in rel if r['id']=='Regular_Lot_3')
    assert regular['interval_seconds']=='3101.4' and regular['wafer_count']==25
    assert regular['due_allowance_us']==4651666542000
    assert all(r['first_at_us']==0 for r in rel)

def test_validation_digest_detects_model_edit(prepared):
    a=deepcopy(prepared);a['model']['routes'][0]['steps'][0]['processing']['mean_seconds']='-1'
    with pytest.raises(ValueError):validate_manufacturing(a)

@pytest.mark.parametrize('sheet,cell,value,match',[
 ('Route_Product_3','S1','PROCESS_OFFSET','header'),
 ('Route_Product_3','I2',999999,'range'),
 ('Route_Product_3','P18','sometimes','WHEN'),
 ('Route_Product_3','X67','9999','rework'),
 ('Route_Product_3','U14','2','dedication'),
 ('PM','F2','bananas','unit'),
 ('PM','C2','missing','scope'),
 ('Breakdown','D2','utilization','type'),
 ('Lotrelease','I2',0,'interval'),
])
def test_required_unsupported_or_invalid_semantics_reject(source,tmp_path,sheet,cell,value,match):
    w=load_workbook(source);w[sheet][cell]=value;p=tmp_path/'mutated.xlsx';w.save(p);w.close()
    with pytest.raises(ValueError,match=match):prepare_manufacturing(p)

def test_process_and_setup_offset_remain_distinct(source,tmp_path):
    w=load_workbook(source);s=w['Route_Product_3']
    for cell,value in [('O2','fixture-setup'),('P2','need'),('Q2','uniform'),('R2',10),('S2',2),('T2','min')]:s[cell]=value
    p=tmp_path/'two-offsets.xlsx';w.save(p);w.close()
    step=route(prepare_manufacturing(p),'Product_3')['steps'][0]
    assert step['processing']['offset_seconds']=='1503.99'
    assert step['setup']['fallback']['offset_seconds']=='120'

@pytest.mark.parametrize('sheet,cell,value,match',[
 ('Route_Product_3','N2',100,'batch'),
 ('Route_Product_3','Z2','99999','cqt'),
 ('Route_Product_3','Y2',101,'probability'),
 ('Setups','F6',2.5,'integer'),
 ('PM','D2','usage based','type'),
 ('PM','H2',10000,'interval'),
 ('Toolgroups','N2','Critical Ratio','ranking'),
 ('Toolgroups','D2','maybe','flag'),
])
def test_additional_required_rules_reject(source,tmp_path,sheet,cell,value,match):
    w=load_workbook(source);w[sheet][cell]=value
    if sheet=='Route_Product_3' and cell=='Z2':
        w[sheet]['AA2']=1;w[sheet]['AB2']='hr'
    p=tmp_path/'invalid.xlsx';w.save(p);w.close()
    with pytest.raises(ValueError,match=match):prepare_manufacturing(p)

@pytest.mark.parametrize('mutation',['delete','swap'])
def test_incomplete_or_reordered_target_route_rejects(source,tmp_path,mutation):
    w=load_workbook(source);s=w['Route_Product_3']
    if mutation=='delete':s.delete_rows(2)
    else:
        a=[c.value for c in s[2]];b=[c.value for c in s[3]]
        for i,(x,y) in enumerate(zip(a,b),1):s.cell(2,i).value=y;s.cell(3,i).value=x
    p=tmp_path/'route.xlsx';w.save(p);w.close()
    with pytest.raises(ValueError,match='complete source sequence'):prepare_manufacturing(p)

def test_partial_setup_fallback_rejects(source,tmp_path):
    w=load_workbook(source);s=w['Route_Product_3'];s['R18']=99;s['T18']='bananas'
    p=tmp_path/'partial-setup.xlsx';w.save(p);w.close()
    with pytest.raises(ValueError,match='partial setup'):prepare_manufacturing(p)

def test_unbounded_calendar_pm_rejects(source,tmp_path):
    w=load_workbook(source);s=w['PM'];s['G2']='exponential';s['I2']=None
    p=tmp_path/'unbounded-pm.xlsx';w.save(p);w.close()
    with pytest.raises(ValueError,match='bounded'):prepare_manufacturing(p)

def test_unsupported_adapter_cannot_pass_self_consistent_binding(prepared):
    from fab_sim.inputs import bind_validation
    from fab_sim.sources.smt import VALIDATOR
    a=deepcopy(prepared);a['provenance']['adapter_version']='unsupported-v999'
    a=bind_validation(a,{'tool_groups':a['model']['required_tool_group_ids']},VALIDATOR)
    with pytest.raises(ValueError,match='adapter'):validate_manufacturing(a)
